#!/usr/bin/env python3
"""SQL-Dokumentation mit Installationsanleitung (_SQL) der Abschlussphase als Word-Datei.

Liest nur Dateien des Repos und schreibt
abgabe/03_Finalisierungsphase/Abschlussphase_SQL.docx:

  1 Installationsanleitung  README.md, Abschnitt „Schnellstart“
  2 Init-Kette              docs/belege/init_log.txt, dann sql/init/0*.sql:
                            handgeschriebene Skripte vollständig, erzeugte
                            (02, 03, 05) mit Kopfkommentar und Verweis auf die
                            Zeile in der SQL-Abgabedatei; zu 06 tests.txt,
                            zu 08 healthcheck_marker.txt
  3 Fachabfragen            sql/queries/0*.sql: je Abfrage Kopfblock, SQL und
                            Ergebnisbeleg aus sql/queries/ergebnis/
  4 Performance             die vier Textpläne (200.000 Zeilen), messreihe_*.md
  5 Nebenläufigkeit         Kopfkommentar von tools/test_nebenlaeufigkeit.sh, Transkript
  6 Metadaten               docs/belege/metadaten.txt

Code und Belege stehen unverändert in Consolas 8 pt, eine Quellzeile je Absatz;
breitere psql-Zeilen bricht Word um. Der Generator schreibt nur Überschriften
und Quellenangaben, keinen Fließtext. Layout wie build_kz.py (Arial, 2 cm Rand,
Kopfzeile Name · Matrikelnummer · Modul · Datum), hier als Seitenkopf auf jeder Seite.

Voraussetzung: python tools/build_sql.py ist gelaufen — die Zeilenverweise für
die erzeugten Init-Skripte stammen aus dem Inhaltsverzeichnis dieser Datei.
Das Inhaltsverzeichnis ist ein Word-Feld; Word füllt es beim PDF-Export
(TablesOfContents(1).Update() per COM), bis dahin steht dort ein Platzhalter.

Aufruf: python tools/build_sqlpdf.py
"""
import datetime
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from build_kz import BLAU, GRAU, SCHRIFT, T, absatz, linie
from build_sql import OUT as SQL_ABGABE, zeilen_lesen
from run_queries import parse_queries

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "abgabe" / "03_Finalisierungsphase" / "Abschlussphase_SQL.docx"

CODE = "Consolas"
# Erzeugt von build_ddl.py / build_seed.py: im PDF nur der Kopfkommentar
# (05_seed.sql allein hat ~1.500 Zeilen), vollständig in der SQL-Abgabedatei.
ERZEUGT = {"02_schema.sql", "03_indexes.sql", "05_seed.sql"}
BELEG_ZU = {"06_tests.sql": "docs/belege/tests.txt",
            "08_init_abschluss.sql": "docs/belege/healthcheck_marker.txt"}
MESSREIHEN = ["messreihe_umkreis.md", "messreihe_trigram.md"]
FELDER = ["entitaet", "frage", "konstrukt", "erwartet", "erklaerung"]


def titel_daten():
    """Name und Matrikelnummer; die Nummer steht nur in der gitignorierten titel_lokal.json."""
    ti = dict(T.get("titel", {}))
    lokal = ROOT / "tools" / "titel_lokal.json"
    if lokal.exists():
        ti.update(json.loads(lokal.read_text(encoding="utf-8")))
    if not all(ti.get(k) for k in ("name", "matrikel")):
        raise SystemExit("texte_phase2.json/titel_lokal.json: name oder matrikel fehlt")
    return ti


def abgabe_name(ti, endung):
    """IU-Dateiname, z. B. Rahman-Tanzin_<Matrikel>_DLBDSPBDM01_D_Abschlussphase_SQL.pdf."""
    teile = ti["name"].split()
    return f"{teile[-1]}-{teile[0]}_{ti['matrikel']}_DLBDSPBDM01_D_Abschlussphase_SQL{endung}"


# ---------------------------------------------------------------------------
# Formatvorlagen
# ---------------------------------------------------------------------------

def schrift(rpr, name):
    """Schrift in rFonts setzen; Theme-Schriften der Vorlage entfernen, sie gingen sonst vor."""
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rf.attrib.pop(qn(f"w:{a}"), None)
    for a in ("ascii", "hAnsi", "eastAsia", "cs"):
        rf.set(qn(f"w:{a}"), name)


def vorlagen(doc):
    st = doc.styles
    schrift(st.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr")), SCHRIFT)
    normal = st["Normal"]
    schrift(normal.element.get_or_add_rPr(), SCHRIFT)
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(4)

    for stufe, groesse in ((1, 14), (2, 12), (3, 11)):
        h = st[f"Heading {stufe}"]
        schrift(h.element.get_or_add_rPr(), SCHRIFT)
        h.font.size = Pt(groesse)
        h.font.bold = True
        h.font.italic = False
        h.font.color.rgb = BLAU
        h.paragraph_format.space_before = Pt(12 if stufe < 3 else 8)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
    # Jedes Kapitel beginnt auf einer neuen Seite.
    st["Heading 1"].paragraph_format.page_break_before = True

    code = st.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    code.base_style = normal
    schrift(code.element.get_or_add_rPr(), CODE)
    code.font.size = Pt(8)
    pf = code.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    pf.widow_control = False
    # Hellgrauer Hintergrund: aufeinanderfolgende Zeilen ergeben einen Block.
    shd = OxmlElement("w:shd")
    for k, v in (("val", "clear"), ("color", "auto"), ("fill", "F2F2F2")):
        shd.set(qn(f"w:{k}"), v)
    code.element.get_or_add_pPr().append(shd)


def feld(p, anweisung, platzhalter):
    """Word-Feld (TOC, PAGE) als begin/instrText/separate/Platzhalter/end."""
    def zeichen(kind):
        p.add_run()._r.append(kind)

    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); zeichen(b)
    i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = anweisung; zeichen(i)
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate"); zeichen(s)
    p.add_run(platzhalter)
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); zeichen(e)


# ---------------------------------------------------------------------------
# Bausteine
# ---------------------------------------------------------------------------

def code(doc, zeilen):
    """Quelltext unverändert, eine Zeile je Absatz; Leerzeilen am Ende entfallen."""
    zeilen = list(zeilen)
    while zeilen and not zeilen[-1].strip():
        zeilen.pop()
    for z in zeilen:
        p = doc.add_paragraph(style="Code")
        if z:
            p.add_run(z.expandtabs(8))
    return len(zeilen)


def quelle(doc, text):
    """Graue Quellenangabe über einem Block, bleibt mit dem Block auf einer Seite."""
    p = absatz(doc, text, size=8.5, color=GRAU, after=2)
    p.paragraph_format.keep_with_next = True
    return p


def inline(p, text, size=None):
    """`Code` in Consolas, **fett** in Fettschrift, sonst Normaltext."""
    for teil in re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text):
        if not teil:
            continue
        if teil.startswith("`"):
            r = p.add_run(teil[1:-1])
            schrift(r._r.get_or_add_rPr(), CODE)
        elif teil.startswith("**"):
            r = p.add_run(teil[2:-2])
            r.bold = True
        else:
            r = p.add_run(teil)
        if size:
            r.font.size = Pt(size)


def tabelle(doc, zeilen, breite_cm=17.0):
    """Markdown-Tabelle (Kopf, Ausrichtungszeile, Daten) als Word-Tabelle in Arial 8.

    Spaltenbreite im Verhältnis zum längsten nicht umbrechbaren Stück der Spalte
    (Zahl wie 1.000.000, Kopfname bis zum nächsten „_“), damit Word keine Zahl
    mitten im Wert umbricht; Kopfnamen dürfen nach „_“ umbrechen.
    """
    zellen = [[c.strip() for c in z.strip().strip("|").split("|")] for z in zeilen]
    kopf, ausrichtung, daten = zellen[0], zellen[1], zellen[2:]
    rechts = [a.endswith(":") for a in ausrichtung]
    bedarf = [max([len(s) for s in kopf[c].split("_")] +
                  [len(s) for zeile in daten for s in zeile[c].split()])
              for c in range(len(kopf))]
    t = doc.add_table(rows=1 + len(daten), cols=len(kopf))
    t.style = doc.styles["Table Grid"]
    t.autofit = False
    # Zellenabstand links/rechts 0,1 cm statt 0,19 cm
    mar = OxmlElement("w:tblCellMar")
    for seite in ("left", "right"):
        m = OxmlElement(f"w:{seite}"); m.set(qn("w:w"), "57"); m.set(qn("w:type"), "dxa")
        mar.append(m)
    look = t._tbl.tblPr.find(qn("w:tblLook"))           # Schema-Reihenfolge: vor tblLook
    if look is not None:
        look.addprevious(mar)
    else:
        t._tbl.tblPr.append(mar)
    for c, b in enumerate(bedarf):
        breite = Cm(breite_cm * b / sum(bedarf))
        t.columns[c].width = breite
        for zelle in t.columns[c].cells:
            zelle.width = breite
    for r, zeile in enumerate([kopf] + daten):
        trpr = t.rows[r]._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))          # Zeile nicht über zwei Seiten
        if r == 0:
            trpr.append(OxmlElement("w:tblHeader"))      # Kopfzeile wiederholen
        for c, wert in enumerate(zeile):
            p = t.cell(r, c).paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if rechts[c] and r > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            # Nullbreites Leerzeichen nach „_“: Umbruchstelle im Kopfnamen
            run = p.add_run(wert.replace("_", "_​") if r == 0 else wert)
            run.font.size = Pt(8)
            run.bold = r == 0
    doc.add_paragraph()


def markdown(doc, zeilen):
    """Kleiner Markdown-Satz für README und Messreihe: Absätze, Codeblöcke, Listen, Tabellen."""
    puffer = []

    def absatz_schliessen():
        if puffer:
            inline(doc.add_paragraph(), " ".join(puffer))
            puffer.clear()

    i = 0
    while i < len(zeilen):
        z = zeilen[i]
        if z.startswith("```"):
            absatz_schliessen()
            ende = next(j for j in range(i + 1, len(zeilen)) if zeilen[j].startswith("```"))
            code(doc, zeilen[i + 1:ende])
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
            i = ende + 1
            continue
        if z.startswith("|"):
            absatz_schliessen()
            ende = i
            while ende < len(zeilen) and zeilen[ende].startswith("|"):
                ende += 1
            tabelle(doc, zeilen[i:ende])
            i = ende
            continue
        if not z.strip():
            absatz_schliessen()
        elif z.startswith("#"):
            absatz_schliessen()
            p = doc.add_paragraph()
            inline(p, "**" + z.lstrip("#").strip() + "**")
        elif z.startswith("- "):
            absatz_schliessen()
            inline(doc.add_paragraph(style="List Bullet"), z[2:], size=9)
        else:
            puffer.append(z.strip())
        i += 1
    absatz_schliessen()


def kopfkommentar(zeilen, zeichen="--"):
    """Führende Kommentarzeilen (ohne Shebang); Rückgabe (erste, letzte Zeilennummer, Zeilen)."""
    start = 1 if zeilen and zeilen[0].startswith("#!") else 0
    ende = start
    while ende < len(zeilen) and zeilen[ende].startswith(zeichen):
        ende += 1
    if ende == start:
        raise SystemExit("kein Kopfkommentar gefunden")
    return start + 1, ende, zeilen[start:ende]


def teil_a_zeilen():
    """{Dateiname: Zeile} der Init-Skripte (Teil A) aus dem Inhaltsverzeichnis der SQL-Abgabe."""
    if not SQL_ABGABE.exists():
        raise SystemExit(f"{SQL_ABGABE.relative_to(ROOT)} fehlt: zuerst python tools/build_sql.py")
    alle = zeilen_lesen(SQL_ABGABE)
    teil, ziel = None, {}
    for z in alle:
        if not z.startswith("--"):
            break
        if m := re.match(r"^-- Teil (\w) ", z):
            teil = m.group(1)
        elif (m := re.match(r"^--   (\S+) \.+ Zeile\s+(\d+)$", z)) and teil == "A":
            ziel[m.group(1)] = int(m.group(2))
    return alle, ziel


# ---------------------------------------------------------------------------
# Kapitel
# ---------------------------------------------------------------------------

def kapitel_installation(doc):
    doc.add_heading("1 Installationsanleitung", level=1)
    zeilen = zeilen_lesen(ROOT / "README.md")
    start = zeilen.index("## Schnellstart") + 1
    ende = next((j for j in range(start, len(zeilen)) if zeilen[j].startswith("## ")), len(zeilen))
    quelle(doc, f"Quelle: README.md, Abschnitt „Schnellstart“ (Zeilen {start + 1}–{ende})")
    markdown(doc, zeilen[start:ende])


def kapitel_init(doc, ti):
    doc.add_heading("2 Init-Kette (sql/init/)", level=1)
    doc.add_heading("docs/belege/init_log.txt", level=2)
    code(doc, zeilen_lesen(ROOT / "docs" / "belege" / "init_log.txt"))

    abgabe, zeile_in_abgabe = teil_a_zeilen()
    sql_name = abgabe_name(ti, ".sql")
    dateien = sorted((ROOT / "sql" / "init").glob("0*.sql"))
    if len(dateien) != 8:
        raise SystemExit(f"sql/init/: {len(dateien)} statt 8 Skripte")
    for f in dateien:
        zeilen = zeilen_lesen(f)
        doc.add_heading(f"sql/init/{f.name}", level=2)
        if f.name in ERZEUGT:
            nr = zeile_in_abgabe.get(f.name)
            # Vier Zeilen Trennkommentar, dann die Datei — sonst ist die SQL-Abgabe veraltet.
            if nr is None or abgabe[nr + 3:nr + 3 + len(zeilen)] != zeilen:
                raise SystemExit(f"{f.name} passt nicht zur SQL-Abgabe: python tools/build_sql.py neu ausführen")
            von, bis, kopf = kopfkommentar(zeilen)
            quelle(doc, f"Kopfkommentar, Zeilen {von}–{bis} von {len(zeilen)} (erzeugte Datei); "
                        f"vollständig in {sql_name}, Teil A, ab Zeile {nr}")
            code(doc, kopf)
        else:
            quelle(doc, f"vollständig, {len(zeilen)} Zeilen")
            code(doc, zeilen)
        if f.name in BELEG_ZU:
            doc.add_heading(BELEG_ZU[f.name], level=3)
            code(doc, zeilen_lesen(ROOT / BELEG_ZU[f.name]))


def kapitel_abfragen(doc):
    doc.add_heading("3 Fachabfragen (sql/queries/)", level=1)
    ergebnis = ROOT / "sql" / "queries" / "ergebnis"
    vorhanden = sorted(p.name for p in ergebnis.glob("[0-9][0-9]_*.txt"))
    nummer, belegt = 0, []
    for f in sorted((ROOT / "sql" / "queries").glob("[0-9][0-9]_*.sql")):
        doc.add_heading(f"sql/queries/{f.name}", level=2)
        zeilen = zeilen_lesen(f)
        erste = next(i for i, z in enumerate(zeilen) if re.match(r"^--\s*@entitaet:", z))
        quelle(doc, f"Dateikopf, Zeilen 1–{erste}")
        code(doc, zeilen[:erste])
        for block in parse_queries(f):
            nummer += 1
            beleg = f"{nummer:02d}_{block['entitaet']}.txt"
            fehlend = [k for k in FELDER if not block.get(k)]
            if fehlend or beleg not in vorhanden:
                raise SystemExit(f"{f.name} Zeile {block['zeile']}: Feld {fehlend} leer oder {beleg} fehlt")
            doc.add_heading(f"Frage {nummer:02d} · {block['entitaet']}", level=3)
            for k in FELDER:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(1)
                inline(p, f"**@{k}:** {block[k]}", size=9)
            quelle(doc, f"SQL: sql/queries/{f.name}, ab Zeile {block['zeile']}")
            code(doc, block["sql"].split("\n"))
            quelle(doc, f"Ergebnis: sql/queries/ergebnis/{beleg}")
            code(doc, zeilen_lesen(ergebnis / beleg))
            belegt.append(beleg)
    if nummer != 21 or sorted(belegt) != vorhanden:
        raise SystemExit(f"{nummer} Abfragen, {len(vorhanden)} Ergebnisdateien — erwartet je 21")
    return nummer


def kapitel_performance(doc):
    doc.add_heading("4 Performance (sql/perf/)", level=1)
    ordner = ROOT / "sql" / "perf" / "ergebnis"
    plaene = sorted(ordner.glob("0*.txt"))
    if len(plaene) != 4:
        raise SystemExit(f"sql/perf/ergebnis/: {len(plaene)} statt 4 Textpläne")
    for f in plaene:
        doc.add_heading(f"sql/perf/ergebnis/{f.name}", level=2)
        code(doc, zeilen_lesen(f))
    for name in MESSREIHEN:
        doc.add_heading(f"sql/perf/ergebnis/{name}", level=2)
        markdown(doc, zeilen_lesen(ordner / name))


def kapitel_nebenlaeufigkeit(doc):
    doc.add_heading("5 Nebenläufigkeit", level=1)
    zeilen = zeilen_lesen(ROOT / "tools" / "test_nebenlaeufigkeit.sh")
    von, bis, kopf = kopfkommentar(zeilen, "#")
    doc.add_heading("tools/test_nebenlaeufigkeit.sh", level=2)
    quelle(doc, f"Kopfkommentar, Zeilen {von}–{bis} von {len(zeilen)}")
    code(doc, kopf)
    doc.add_heading("docs/belege/nebenlaeufigkeit.txt", level=2)
    code(doc, zeilen_lesen(ROOT / "docs" / "belege" / "nebenlaeufigkeit.txt"))


def kapitel_metadaten(doc):
    doc.add_heading("6 Metadaten", level=1)
    doc.add_heading("docs/belege/metadaten.txt", level=2)
    code(doc, zeilen_lesen(ROOT / "docs" / "belege" / "metadaten.txt"))


def main():
    ti = titel_daten()
    datum = datetime.date.today().strftime("%d.%m.%Y")

    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for seite in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, seite, Cm(2.0))
    vorlagen(doc)
    doc.core_properties.author = ti["name"]
    doc.core_properties.title = "SQL-Dokumentation mit Installationsanleitung"

    # Kopfzeile wie build_kz.py, hier im Seitenkopf: Nachname, Vorname(n) · Matrikelnummer · Modul · Datum
    teile = ti["name"].split()
    kopfname = f"{teile[-1]}, {' '.join(teile[:-1])}" if len(teile) > 1 else ti["name"]
    kp = sec.header.paragraphs[0]
    kp.paragraph_format.space_after = Pt(2)
    r = kp.add_run(f"{kopfname}  ·  Matrikelnummer {ti['matrikel']}  ·  DLBDSPBDM01_D  ·  {datum}")
    r.font.size = Pt(8.5)
    r.font.color.rgb = GRAU
    linie(kp)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    feld(fp, "PAGE", "1")

    absatz(doc, "Portfolio – Abschlussphase", size=14, bold=True, color=BLAU, after=3)
    absatz(doc, "Data-Mart-Erstellung in SQL: Datenbank für eine Buchtausch-App", size=12, color=GRAU, after=9)
    absatz(doc, "SQL-Dokumentation mit Installationsanleitung", size=12, bold=True, after=12)
    absatz(doc, "Inhalt", size=12, bold=True, color=BLAU, after=6)
    feld(doc.add_paragraph(), r'TOC \o "1-3" \h \z \u',
         "Inhaltsverzeichnis: wird beim PDF-Export von Word aktualisiert.")

    kapitel_installation(doc)
    kapitel_init(doc, ti)
    abfragen = kapitel_abfragen(doc)
    kapitel_performance(doc)
    kapitel_nebenlaeufigkeit(doc)
    kapitel_metadaten(doc)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    ueberschriften = [p.style.name for p in doc.paragraphs if p.style.name.startswith("Heading")]
    codezeilen = sum(1 for p in doc.paragraphs if p.style.name == "Code")
    print(f"geschrieben: {OUT.relative_to(ROOT)}  ({ueberschriften.count('Heading 1')} Kapitel, "
          f"{len(ueberschriften)} Überschriften, {abfragen} Abfragen mit Ergebnis, {codezeilen} Codezeilen)")


if __name__ == "__main__":
    main()
