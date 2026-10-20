#!/usr/bin/env python3
"""
Baut die Präsentation der Phase 2 (Erarbeitungs-/Reflexionsphase) als .pptx.

Aufruf:   python3 tools/build_slides.py            (von überall; python-pptx + Pillow)
Ausgabe:  abgabe/02_Erarbeitungsphase/Erarbeitungsphase_PR.pptx
          docs/screenshots/*.png   (Ergebnis-„Screenshots“ in psql-Optik)

Quellen — nichts wird hier von Hand gepflegt:
  tools/model.py                 Entitäten, Gruppen, Farben, Indizes
  sql/init/02_schema.sql         CREATE TABLE je Entität (dokumentiertes SQL-Statement)
  sql/init/04_regeln.sql         Trigger-Funktionen (Umsetzung der Regeln 2–4)
  sql/init/06_tests.sql          Negativtest-Statements je Regel
  sql/queries/*.sql              Testfall je Entität (Kopfblock @entitaet …), via run_queries.parse_queries()
  sql/queries/ergebnis/*.txt     Ergebnis des Testfalls (→ Screenshot)
  sql/perf/ergebnis/*.txt        EXPLAIN-Pläne mit/ohne Index
  docs/belege/init_log.txt       Auszug aus `docker compose logs db` (Reproduzierbarkeit)
  docs/belege/tests.txt          Ergebnistabelle aus 06_tests.sql (aus dem Container)
  docs/belege/zeilenzahlen.txt   Zeilen je Tabelle nach frischem Start
  docs/erm/ER-Modell_Buchtausch.png
  tools/texte_phase2.json        alle Folientexte (schreibt Tanzin); @erklaerung je Entität
                                 steht direkt im Kopfblock der Abfrage in sql/queries/
  tools/titel_lokal.json         Matrikelnummer (gitignoriert, ergänzt "titel")

Platzhalter: Texte, die noch fehlen, erscheinen orange als „[Tanzin: …]“, damit
nichts Unfertiges unbemerkt in die Abgabe rutscht.
"""
import json
import re
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

TOOLS = Path(__file__).resolve().parent
REPO = TOOLS.parent
sys.path.insert(0, str(TOOLS))
from model import E, R, GROUPS, TERNAER, INDEXES          # noqa: E402
from run_queries import parse_queries                    # noqa: E402

OUT_DIR = REPO / "abgabe" / "02_Erarbeitungsphase"
OUT = OUT_DIR / "Erarbeitungsphase_PR.pptx"
SHOTS = REPO / "docs" / "screenshots"
TEXTE = json.loads((TOOLS / "texte_phase2.json").read_text(encoding="utf-8"))

# ------------------------------------------------------------------ Gestaltung
W, H = 13.333, 7.5
DUNKEL = RGBColor(0x23, 0x32, 0x4A)      # Titel-/Schlussfolien, Überschriften
GRAU = RGBColor(0x5F, 0x66, 0x73)        # Nebentext
HELL = RGBColor(0xF3, 0xF4, 0xF6)        # Codekästen
KOMMENTAR = RGBColor(0x6B, 0x8E, 0x23)   # SQL-Kommentare im Code
ORANGE = RGBColor(0xC2, 0x5A, 0x00)      # Platzhalter
WEISS = RGBColor(0xFF, 0xFF, 0xFF)
SCHRIFT = "Calibri"
CODE = "Courier New"
MONO_TTF = next((p for p in [
    "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf",
    "C:/Windows/Fonts/consola.ttf", "C:/Windows/Fonts/cour.ttf",
    "/System/Library/Fonts/Menlo.ttc"] if Path(p).exists()), None)


def rgb(hexstr):
    return RGBColor.from_string(hexstr.lstrip("#"))


def platzhalter(text):
    return text is None or text.strip() in ("", "(Tanzin)") or text.strip().startswith("[Tanzin")


# ------------------------------------------------------------------ Textbausteine
def box(slide, x, y, w, h, text="", size=14, bold=False, color=DUNKEL, font=SCHRIFT,
        align=PP_ALIGN.LEFT, italic=False, anchor=MSO_ANCHOR.TOP, margin=0.05):
    """Einfaches Textfeld; `text` darf Zeilenumbrüche enthalten (je Absatz)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, Inches(margin))
    zeilen = text.split("\n") if text else [""]
    for i, z in enumerate(zeilen):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = z
        f = r.font
        f.name, f.size, f.bold, f.italic = font, Pt(size), bold, italic
        f.color.rgb = color
    return tb


def text_oder_platzhalter(slide, x, y, w, h, text, hinweis, size=14, **kw):
    if platzhalter(text):
        kw.update(color=ORANGE, italic=True)
        return box(slide, x, y, w, h, f"[Tanzin: {hinweis}]", size=size, **kw)
    return box(slide, x, y, w, h, text, size=size, **kw)


def absaetze(slide, x, y, w, h, texte, hinweis, size=14, abstand=6):
    """Mehrere Absätze untereinander; Liste leerer Texte -> Platzhalter."""
    if not texte or all(platzhalter(t) for t in texte):
        return box(slide, x, y, w, h, f"[Tanzin: {hinweis}]", size=size, color=ORANGE, italic=True)
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, t in enumerate(texte):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(abstand)
        r = p.add_run()
        r.text = t if not platzhalter(t) else "[Tanzin: Absatz fehlt]"
        r.font.name, r.font.size = SCHRIFT, Pt(size)
        r.font.color.rgb = ORANGE if platzhalter(t) else DUNKEL
        r.font.italic = platzhalter(t)
    return tb


def aufzaehlung(slide, x, y, w, h, punkte, hinweis, size=14):
    if not punkte or all(platzhalter(t) for t in punkte):
        return box(slide, x, y, w, h, f"[Tanzin: {hinweis}]", size=size, color=ORANGE, italic=True)
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, t in enumerate(punkte):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        r = p.add_run()
        r.text = "–  " + (t if not platzhalter(t) else "[Tanzin: Punkt fehlt]")
        r.font.name, r.font.size = SCHRIFT, Pt(size)
        r.font.color.rgb = ORANGE if platzhalter(t) else DUNKEL
    return tb


def titel(slide, text, untertitel=None, chip=None, chip_farben=None):
    box(slide, 0.5, 0.3, 10.6, 0.7, text, size=26 if len(text) <= 48 else 21, bold=True)
    if untertitel is not None:
        text_oder_platzhalter(slide, 0.5, 0.95, 12.3, 0.55, untertitel,
                              "ein Satz, was diese Folie zeigt (@erklaerung)", size=13, italic=True, color=GRAU)
    if chip:
        fill, line = chip_farben
        s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.2), Inches(0.38), Inches(1.65), Inches(0.42))
        s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
        s.line.color.rgb = rgb(line); s.line.width = Pt(1)
        tf = s.text_frame; tf.margin_left = tf.margin_right = Inches(0.05)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = chip
        r.font.name, r.font.size, r.font.bold = SCHRIFT, Pt(10), True
        r.font.color.rgb = rgb(line)


def fusszeile(slide, nr, links=""):
    box(slide, 0.5, 7.05, 10.5, 0.3, links, size=9, color=GRAU)
    box(slide, 12.0, 7.05, 0.85, 0.3, str(nr), size=9, color=GRAU, align=PP_ALIGN.RIGHT)


# ------------------------------------------------------------------ Code
def umbrechen(zeile, breite):
    """Weicher Umbruch langer Codezeilen; Folgezeilen eingerückt."""
    if len(zeile) <= breite:
        return [zeile]
    einzug = len(zeile) - len(zeile.lstrip()) + 4
    teile, rest = [], zeile
    while len(rest) > breite:
        schnitt = rest.rfind(" ", 0, breite)
        if schnitt <= einzug:
            schnitt = breite
        teile.append(rest[:schnitt].rstrip())
        rest = " " * einzug + rest[schnitt:].lstrip()
    teile.append(rest)
    return teile


def code_box(slide, x, y, w, h, zeilen, size=8.5, label=None, h_max=None):
    """Grauer Kasten mit Monospace-Text. Kommentare (-- …) grün. Passt die
    Schriftgröße an, bis der Text in die Höhe h passt (min. 6,5 pt); reicht das
    nicht, darf der Kasten bis h_max wachsen (Standard: h). Rückgabe: Kasten,
    tatsächliche Höhe, Schriftgröße — Folgeelemente richten sich nach der Höhe."""
    h_max = h if h_max is None else h_max
    if label:
        box(slide, x, y - 0.28, w, 0.28, label, size=10, bold=True, color=GRAU)
    while True:
        breite = int((w - 0.2) / (size * 0.6 / 72))       # Zeichen je Zeile
        flach = []
        for z in zeilen:
            flach += umbrechen(z.rstrip(), breite)
        hoehe = len(flach) * size * 1.22 / 72 + 0.15
        if hoehe <= h or size <= 6.5:
            break
        size -= 0.5
    if hoehe > h and h_max > h:
        h = min(h_max, hoehe)        # Kasten wachsen lassen statt Text überlaufen
    rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(min(h, max(hoehe, 0.4))))
    rect.adjustments[0] = 0.03
    rect.fill.solid(); rect.fill.fore_color.rgb = HELL
    rect.line.fill.background()
    tf = rect.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = Inches(0.1)
    tf.margin_top = tf.margin_bottom = Inches(0.06)
    for i, z in enumerate(flach):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(0)
        p.alignment = PP_ALIGN.LEFT
        code, kommentar = z, None
        m = re.match(r"^(.*?)(\s*--.*)$", z)
        if m and not z.lstrip().startswith("--"):
            code, kommentar = m.group(1), m.group(2)
        elif z.lstrip().startswith("--"):
            code, kommentar = "", z
        for txt, farbe in ((code, DUNKEL), (kommentar, KOMMENTAR)):
            if txt:
                r = p.add_run(); r.text = txt
                r.font.name, r.font.size = CODE, Pt(size)
                r.font.color.rgb = farbe
    return rect, min(h, max(hoehe, 0.4)), size


# ------------------------------------------------------------------ Terminal-Bilder
def terminal_png(zeilen, name, kopf="psql (16.x) — buchtausch@localhost:5433"):
    """Rendert Textzeilen als Bild in psql-/Terminal-Optik und gibt (Pfad, Breite, Höhe) zurück."""
    SHOTS.mkdir(parents=True, exist_ok=True)
    fs = 15
    font = ImageFont.truetype(MONO_TTF, fs) if MONO_TTF else ImageFont.load_default()
    cw = font.getlength("M") if hasattr(font, "getlength") else 9
    lh = int(fs * 1.4)
    pad = 16
    breite = int(max(len(z) for z in zeilen + [kopf]) * cw) + 2 * pad
    hoehe = lh * (len(zeilen) + 1) + 2 * pad + 8
    img = Image.new("RGB", (breite, hoehe), (0x1F, 0x24, 0x30))
    d = ImageDraw.Draw(img)
    d.rectangle([0, 0, breite, lh + pad], fill=(0x2B, 0x33, 0x44))
    d.text((pad, pad // 2 + 2), kopf, font=font, fill=(0xB0, 0xB8, 0xC4))
    y = lh + pad + 6
    for z in zeilen:
        farbe = (0xE8, 0xEA, 0xED)
        if z.startswith("buchtausch=#"):
            farbe = (0x8F, 0xD3, 0xA8)
        elif z.startswith("ERROR") or "NICHT bestanden" in z or " FEHL" in z:
            farbe = (0xF2, 0x8B, 0x82)
        elif z.startswith("NOTICE") or "bestanden" in z:
            farbe = (0x9E, 0xD6, 0x7C)
        elif z.startswith("--"):
            farbe = (0x8A, 0x93, 0xA3)
        d.text((pad, y), z, font=font, fill=farbe)
        y += lh
    pfad = SHOTS / f"{name}.png"
    img.save(pfad)
    return pfad, breite, hoehe


def bild_einpassen(slide, pfad, px_w, px_h, x, y, w, h, dpi=105):
    """Bild proportional in den Kasten (x, y, w, h) einpassen, oben links ausgerichtet."""
    nat_w, nat_h = px_w / dpi, px_h / dpi
    s = min(w / nat_w, h / nat_h, 1.0 if dpi > 1 else 10.0)
    return slide.shapes.add_picture(str(pfad), Inches(x), Inches(y), Inches(nat_w * s), Inches(nat_h * s))


def ergebnis_zeilen(pfad, max_zeilen=14):
    """Belegdatei -> Terminalzeilen: Kopfkommentare und SET-Zeile weg, lange Ergebnisse kürzen."""
    roh = [z.rstrip() for z in Path(pfad).read_text(encoding="utf-8").splitlines()]
    roh = [z for z in roh if not z.startswith("-- ") and z.strip() not in ("", "SET")]
    ende = next((i for i, z in enumerate(roh) if re.match(r"^\(\d+ rows?\)", z)), None)
    if ende is not None:
        roh = roh[:ende + 1]          # nur das erste Ergebnis (weitere Statements: siehe Datei)
    if len(roh) <= max_zeilen + 3:
        return roh
    # Tabellenkopf (2 Zeilen) + max_zeilen Datenzeilen + Hinweis + „(n rows)“
    kopf, rest = roh[:2], roh[2:]
    daten = [z for z in rest if not re.match(r"^\(\d+ rows?\)", z)]
    anzahl = [z for z in rest if re.match(r"^\(\d+ rows?\)", z)]
    return kopf + daten[:max_zeilen] + [f"... {len(daten) - max_zeilen} weitere Zeilen, vollständig in {Path(pfad).name}"] + anzahl[:1]


def erstes_statement(zeilen):
    """Nur das erste SQL-Statement (bis zum ersten ';' außerhalb eines Kommentars);
    folgen weitere (Sicht-Variante, EXPLAIN-Demo), steht ein Hinweis darunter."""
    while zeilen and not zeilen[0].strip():
        zeilen.pop(0)
    out = []
    for i, z in enumerate(zeilen):
        out.append(z)
        code = z.split("--")[0]
        if ";" in code:
            rest = [r for r in zeilen[i + 1:] if r.strip() and not r.strip().startswith("--")]
            if rest:
                out.append("-- … weitere Statements dieses Testfalls (z. B. View/EXPLAIN-Demo) in der Datei")
            break
    return out


def kompakte_testtabelle(zeilen, regel=None, breite_fall=52, breite_erg=62):
    """psql-Tabelle aus 06_tests.sql in eine schmalere Darstellung umbauen."""
    out = [f"{'nr':>2}  {'R':<2} {'Testfall':<{breite_fall}} {'ok':<5} Ergebnis"]
    for z in zeilen:
        teile = [t.strip() for t in z.split("|")]
        if len(teile) < 6 or not teile[0].isdigit():
            continue
        nr_, r, fall, _erw, erg, ok = teile[:6]
        if regel and r != regel:
            continue
        fall = fall if len(fall) <= breite_fall else fall[:breite_fall - 1] + "…"
        erg = erg if len(erg) <= breite_erg else erg[:breite_erg - 1] + "…"
        out.append(f"{nr_:>2}  {r:<2} {fall:<{breite_fall}} {('ok' if ok == 't' else 'FEHL'):<5} {erg}")
    return out


def testtabelle_zweizeilig(zeilen, regel, breite=66):
    """Je Testfall zwei Zeilen: Fall + Häkchen, darunter die Datenbankantwort."""
    out = []
    for z in zeilen:
        teile = [t.strip() for t in z.split("|")]
        if len(teile) < 6 or not teile[0].isdigit() or teile[1] != regel:
            continue
        nr_, _r, fall, _erw, erg, ok = teile[:6]
        out.append(f"{nr_:>2}  {fall:<{breite - 6}} {'ok' if ok == 't' else 'FEHL'}")
        erg = erg if len(erg) <= breite - 6 else erg[:breite - 7] + "…"
        out.append(f"      → {erg}")
    return out


def zeilenzahlen_kompakt(zeilen, spalten=3):
    """'tabelle | n'-Zeilen in mehrere Spalten legen."""
    paare = []
    for z in zeilen:
        t = [x.strip() for x in z.split("|")]
        if len(t) == 2 and t[1].isdigit():
            paare.append((t[0], t[1]))
    if not paare:
        return zeilen
    n = (len(paare) + spalten - 1) // spalten
    out = [z for z in zeilen if z.startswith("buchtausch=#") or z.startswith("--")]
    for i in range(n):
        teile = []
        for c in range(spalten):
            j = i + c * n
            if j < len(paare):
                teile.append(f"{paare[j][0]:<18} {paare[j][1]:>5}")
        out.append("   ".join(teile))
    out.append(f"({len(paare)} Tabellen, Summe {sum(int(v) for _, v in paare)} Zeilen)")
    return out


# ------------------------------------------------------------------ Quellen lesen
SCHEMA = (REPO / "sql/init/02_schema.sql").read_text(encoding="utf-8").splitlines()
REGELN_SQL = (REPO / "sql/init/04_regeln.sql").read_text(encoding="utf-8")
TESTS_SQL = (REPO / "sql/init/06_tests.sql").read_text(encoding="utf-8")


def ddl(entitaet):
    """CREATE-TABLE-Block einer Entität aus 02_schema.sql (ohne COMMENT ON)."""
    start = next(i for i, z in enumerate(SCHEMA) if z.startswith(f"CREATE TABLE {entitaet} ("))
    ende = next(i for i in range(start, len(SCHEMA)) if SCHEMA[i].startswith(");"))
    return SCHEMA[start:ende + 1]


def trigger_funktion(name):
    m = re.search(rf"(CREATE OR REPLACE FUNCTION {name}\(\).*?\n\$\$;)", REGELN_SQL, re.S)
    return m.group(1).splitlines()


def negativtests():
    """Aus 06_tests.sql: {Testfall-Label: Statement, das absichtlich scheitert}."""
    out = {}
    for block in re.findall(r"\n    BEGIN\n(.*?)\n    EXCEPTION", TESTS_SQL, re.S):
        stmt = block.split("INSERT INTO testergebnis")[0].strip()
        m = re.search(r"VALUES \('(\d)', '([^']+)'", block)
        if m:
            out[m.group(2)] = (m.group(1), [z[8:] if z.startswith("        ") else z for z in stmt.splitlines()])
    return out


def beleg(name):
    p = REPO / "docs/belege" / name
    if not p.exists():
        return [f"[Beleg fehlt: docs/belege/{name}]"]
    return [z.rstrip() for z in p.read_text(encoding="utf-8").splitlines() if z.strip()]


def plan_zeilen(name):
    z = [l.rstrip() for l in (REPO / "sql/perf/ergebnis" / name).read_text(encoding="utf-8").splitlines()]
    z = [l for l in z if not l.startswith("--") and l.strip() and not set(l.strip()) <= set("-+")]
    # nur den Plan: ab der ersten Zeile mit "(cost="
    start = next((i for i, l in enumerate(z) if "(cost=" in l), 0)
    z = [re.sub(r"\s*\(cost=[^)]*\)", "", l) for l in z[start:]]
    # Geografie-Literal (langer Hex-String) durch einen Platzhalter ersetzen
    return [re.sub(r"'0101000020E6[0-9A-F]+'::geography", "'<Punkt Nordend>'::geography", l) for l in z]


def ausfuehrungszeit(name):
    m = re.search(r"Execution Time: ([\d.]+) ms", (REPO / "sql/perf/ergebnis" / name).read_text(encoding="utf-8"))
    return float(m.group(1)) if m else None


# ------------------------------------------------------------------ Folien
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
LEER = prs.slide_layouts[6]
nr = 0


def neue_folie(dunkel=False):
    global nr
    nr += 1
    s = prs.slides.add_slide(LEER)
    if dunkel:
        bg = s.background.fill
        bg.solid(); bg.fore_color.rgb = DUNKEL
    return s


T = TEXTE

# 1 Titel
s = neue_folie(dunkel=True)
box(s, 0.8, 1.6, 11.7, 1.2, "Buchtausch-App — Datenbankimplementierung", size=40, bold=True, color=WEISS)
box(s, 0.8, 2.8, 11.7, 0.8, "Erarbeitungs-/Reflexionsphase · Portfolio „Data-Mart-Erstellung in SQL“ (DLBDSPBDM01_D)",
    size=18, color=RGBColor(0xC9, 0xD3, 0xE3))
ti = dict(T.get("titel", {}))
# Matrikelnummer steht nicht im Repo, sondern in der gitignorierten tools/titel_lokal.json
if (TOOLS / "titel_lokal.json").exists():
    ti.update(json.loads((TOOLS / "titel_lokal.json").read_text(encoding="utf-8")))
tv = lambda k, h: ti.get(k) or f"[Tanzin: {h}]"
box(s, 0.8, 4.6, 11.7, 1.4,
    f"{tv('name', 'Name')}  ·  Matrikelnummer {tv('matrikel', 'Matrikelnummer')}\n"
    f"{tv('datum', 'Datum')}  ·  IU Internationale Hochschule",
    size=16, color=RGBColor(0xC9, 0xD3, 0xE3))
box(s, 0.8, 6.4, 11.7, 0.5, "PostgreSQL 16 · PostGIS 3.4 · Docker Compose · 21 Entitäten · 3 Dreifachbeziehungen",
    size=12, color=RGBColor(0x9A, 0xA8, 0xBF))

# 2 Agenda
s = neue_folie()
titel(s, "Agenda")
punkte = ["Ausgangslage: Feedback aus Phase 1 und die drei Schwerpunkte",
          "ER-Modell und Werkzeugkette (model.py als einzige Quelle der Wahrheit)",
          "Reproduzierbarkeit: Docker Compose und die Init-Kette",
          "21 Entitäten: SQL-Statement, Testfall, Ergebnis",
          "Vier Geschäftsregeln mit Negativtests",
          "Performance: Indizes mit EXPLAIN ANALYZE belegt",
          "Zusammenfassung der Implementierung und Reflexion"]
for i, p in enumerate(punkte):
    y = 1.4 + i * 0.72
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.7), Inches(y), Inches(0.5), Inches(0.5))
    c.fill.solid(); c.fill.fore_color.rgb = DUNKEL; c.line.fill.background()
    tf = c.text_frame; tf.margin_left = tf.margin_right = 0
    pp = tf.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER
    r = pp.add_run(); r.text = str(i + 1); r.font.size = Pt(14); r.font.bold = True; r.font.color.rgb = WEISS; r.font.name = SCHRIFT
    box(s, 1.45, y + 0.02, 10.5, 0.5, p, size=18, anchor=MSO_ANCHOR.MIDDLE)
fusszeile(s, nr)

# 3 Ausgangslage
s = neue_folie()
titel(s, "Ausgangslage: Feedback aus Phase 1")
karten = [("1", "Geschäftsregeln", "auf Datenbankebene, je mit Negativtest"),
          ("2", "Reproduzierbarkeit", "docker compose down -v; up -d --build"),
          ("3", "Performance", "Indizes begründet und mit EXPLAIN ANALYZE belegt")]
aus = T.get("ausgangslage", [])
for i, (n, kopf, sub) in enumerate(karten):
    x = 0.6 + i * 4.15
    k = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.5), Inches(3.9), Inches(4.9))
    k.adjustments[0] = 0.04; k.fill.solid(); k.fill.fore_color.rgb = HELL; k.line.fill.background()
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.3), Inches(1.8), Inches(0.6), Inches(0.6))
    c.fill.solid(); c.fill.fore_color.rgb = DUNKEL; c.line.fill.background()
    r = c.text_frame.paragraphs[0]; r.alignment = PP_ALIGN.CENTER
    rr = r.add_run(); rr.text = n; rr.font.size = Pt(16); rr.font.bold = True; rr.font.color.rgb = WEISS
    box(s, x + 1.0, 1.8, 2.7, 0.6, kopf, size=18, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    box(s, x + 0.3, 2.6, 3.3, 0.7, sub, size=12, color=GRAU, italic=True)
    text_oder_platzhalter(s, x + 0.3, 3.4, 3.3, 2.8, aus[i] if i < len(aus) else None,
                          "2–3 Sätze: was der Tutor erwartet und wie es umgesetzt wurde", size=13)
fusszeile(s, nr, "Tutor-Feedback Phase 1 (17.09.2026): Regeln mit Negativtests, Reproduzierbarkeit, Performance")

# 4 ER-Modell
s = neue_folie()
titel(s, "ER-Modell (Krähenfuß-Notation)", T.get("erm"))
with Image.open(REPO / "docs/erm/ER-Modell_Buchtausch.png") as im:
    bild_einpassen(s, REPO / "docs/erm/ER-Modell_Buchtausch.png", im.width, im.height, 0.5, 1.55, 9.3, 5.35, dpi=1)
stats = [(str(len(E)), "Entitäten"), (str(len(R)), "Beziehungen"),
         (str(sum(len(a) for _g, _z, a in E.values())), "Attribute"), (str(len(TERNAER)), "Dreifach-\nbeziehungen")]
for i, (zahl, lab) in enumerate(stats):
    y = 1.6 + i * 1.3
    box(s, 10.2, y, 2.6, 0.8, zahl, size=40, bold=True, color=DUNKEL)
    box(s, 10.2, y + 0.75, 2.6, 0.5, lab, size=12, color=GRAU)
fusszeile(s, nr, "Erzeugt aus tools/model.py mit build_erm.py (Graphviz) — unverändert seit Phase 1")

# 5 Werkzeugkette
s = neue_folie()
titel(s, "Werkzeugkette: eine Quelle, alle Artefakte", None)
absaetze(s, 0.5, 0.95, 12.3, 0.9, T.get("werkzeugkette", []), "2 Sätze: warum model.py die einzige Quelle ist und was daraus entsteht", size=13)
q = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(3.2), Inches(2.4), Inches(1.4))
q.adjustments[0] = 0.08; q.fill.solid(); q.fill.fore_color.rgb = DUNKEL; q.line.fill.background()
tf = q.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "tools/model.py"; r.font.name = CODE; r.font.size = Pt(14); r.font.bold = True; r.font.color.rgb = WEISS
p = tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "21 Entitäten · Constraints · Indizes"; r.font.size = Pt(10); r.font.color.rgb = RGBColor(0xC9, 0xD3, 0xE3)
ketten = [("build_ddl.py", "02_schema.sql · 03_indexes.sql"),
          ("build_seed.py", "05_seed.sql (Faker, Seed 42)"),
          ("build_erm.py", "ER-Diagramm (Graphviz)"),
          ("build_dict.py", "Datenwörterbuch"),
          ("build_slides.py", "diese Präsentation")]
for i, (werk, art) in enumerate(ketten):
    y = 2.0 + i * 0.78
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(3.2), Inches(y + 0.15), Inches(0.7), Inches(0.3))
    a.fill.solid(); a.fill.fore_color.rgb = GRAU; a.line.fill.background()
    b1 = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.1), Inches(y), Inches(2.5), Inches(0.6))
    b1.adjustments[0] = 0.15; b1.fill.solid(); b1.fill.fore_color.rgb = HELL; b1.line.fill.background()
    t1 = b1.text_frame; t1.vertical_anchor = MSO_ANCHOR.MIDDLE; pp = t1.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER
    rr = pp.add_run(); rr.text = werk; rr.font.name = CODE; rr.font.size = Pt(12); rr.font.color.rgb = DUNKEL
    a2 = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(6.8), Inches(y + 0.15), Inches(0.7), Inches(0.3))
    a2.fill.solid(); a2.fill.fore_color.rgb = GRAU; a2.line.fill.background()
    box(s, 7.7, y, 5.1, 0.6, art, size=13, anchor=MSO_ANCHOR.MIDDLE)
box(s, 0.5, 6.15, 12.3, 0.8,
    "Stack: PostgreSQL 16.4 + PostGIS 3.4.3 im Docker-Image buchtausch-db (Dockerfile auf postgis/postgis:16-3.4) · "
    "Python 3 + Faker für den Seed · Graphviz · psql · DBeaver · Git",
    size=11, color=GRAU)
fusszeile(s, nr, "Regel: Modelländerung nur in model.py, danach alle Artefakte neu erzeugen — nie einzeln nachpflegen")

# 6 Reproduzierbarkeit
s = neue_folie()
titel(s, "Reproduzierbarkeit: frischer Start in einem Befehl", None)
absaetze(s, 0.5, 0.95, 6.0, 1.6, T.get("reproduzierbarkeit", []), "3 Sätze: Init-Kette im Image, Reset-Befehl, was das Log belegt", size=11.5)  # 94 Wörter: bei 12 pt 0,05 in Überlauf (Messung 13.10.)
kette = ["01_extensions.sql   postgis, btree_gist, pg_trgm",
         "02_schema.sql       21 Tabellen, PK/FK/UNIQUE/CHECK, EXCLUDE   (generiert)",
         "03_indexes.sql      20 FK-Indizes + 7 fachliche Indizes        (generiert)",
         "04_regeln.sql       4 Trigger (Regeln 2–4)",
         "05_seed.sql         1.357 Zeilen, Seed 42, Stichtag 19.09.2026 (generiert)",
         "06_tests.sql        16 Regeltests, bricht den Start bei Fehler ab",
         "07_views.sql        View v_exemplar_suche",
         "",
         "$ docker compose down -v; docker compose up -d --build"]
code_box(s, 0.5, 2.9, 6.0, 2.6, kette, size=9, label="Init-Kette im Image (Dockerfile: COPY sql/init/)")
log_zeilen = []
for z in beleg("init_log.txt")[:14]:
    z = z.replace("/usr/local/bin/docker-entrypoint.sh: ", "")
    log_zeilen += umbrechen(z, 84)      # lange Zeilen (NOTICE, LOG) umbrechen statt die Schrift zu verkleinern
pfad, pw, ph = terminal_png(log_zeilen, "repro_log", kopf="docker compose logs db")
bild_einpassen(s, pfad, pw, ph, 6.8, 1.0, 6.0, 3.0)
pfad, pw, ph = terminal_png(zeilenzahlen_kompakt(beleg("zeilenzahlen.txt")), "repro_zeilen")
bild_einpassen(s, pfad, pw, ph, 6.8, 4.2, 6.0, 2.7)
fusszeile(s, nr, "Healthcheck prüft über TCP (pg_isready -h localhost): „healthy“ erst nach Abschluss der Init-Kette")

# 7–27 Entitäten
bloecke = []
for f in sorted((REPO / "sql/queries").glob("0*.sql")):
    bloecke += parse_queries(f)
nach_entitaet = {b["entitaet"]: b for b in bloecke}
belege = {p.name.split("_", 1)[1][:-4]: p for p in sorted((REPO / "sql/queries/ergebnis").glob("*.txt"))}
lauf = 0
for name, (grp, zweck, attrs) in E.items():
    lauf += 1
    b = nach_entitaet.get(name)
    if b is None:
        raise SystemExit(f"Keine Abfrage für Entität {name} in sql/queries/")
    s = neue_folie()
    gtitel, gfill, gline = GROUPS[grp]
    kopf = f"{name}" + ("  ⟨Dreifachbeziehung⟩" if name in TERNAER else "")
    titel(s, kopf, b.get("erklaerung"), chip=gtitel, chip_farben=(gfill, gline))
    # Testfall-SQL ohne Kopfblock
    sql_zeilen = erstes_statement([z for z in b["sql"].splitlines() if not z.startswith("-- @")])
    erg = ergebnis_zeilen(belege[name])
    breit = max(len(z) for z in erg) > 72
    if breit:
        # Layout B: DDL links, Testfall rechts, Ergebnis in voller Breite unten
        _, h1, _ = code_box(s, 0.5, 1.85, 6.0, 2.85, ddl(name), size=8, label="SQL-Statement (sql/init/02_schema.sql)", h_max=3.6)
        _, h2, _ = code_box(s, 6.8, 1.85, 6.0, 2.85, sql_zeilen, size=8, label=f"Testfall ({b['datei']})", h_max=3.6)
        y_bild = 1.85 + max(h1, h2) + 0.35
        if 7.0 - y_bild < 2.4:
            # wenig Platz unter hohen Codekästen: weniger Zeilen zeigen, damit das Bild lesbar bleibt
            erg = ergebnis_zeilen(belege[name], max_zeilen=max(6, int((7.0 - y_bild) * 5)))
        box(s, 0.5, y_bild - 0.28, 6, 0.28, "Ergebnis (psql)", size=10, bold=True, color=GRAU)
        pfad, pw, ph = terminal_png(erg, f"{lauf:02d}_{name}")
        bild_einpassen(s, pfad, pw, ph, 0.5, y_bild, 12.3, 7.0 - y_bild)
    else:
        # Layout A: DDL links über volle Höhe, rechts Testfall und Ergebnis
        code_box(s, 0.5, 1.85, 6.0, 5.1, ddl(name), size=8.5, label="SQL-Statement (sql/init/02_schema.sql)")
        _, h2, _ = code_box(s, 6.8, 1.85, 6.0, 2.9, sql_zeilen, size=8.5, label=f"Testfall ({b['datei']})", h_max=3.6)
        y_bild = 1.85 + h2 + 0.4
        box(s, 6.8, y_bild - 0.28, 6, 0.28, "Ergebnis (psql)", size=10, bold=True, color=GRAU)
        pfad, pw, ph = terminal_png(erg, f"{lauf:02d}_{name}")
        bild_einpassen(s, pfad, pw, ph, 6.8, y_bild, 6.0, 7.0 - y_bild)
    fusszeile(s, nr, f"Konstrukte: {b.get('konstrukt', '')}   ·   Erwartet: {b.get('erwartet', '')}"[:170])

# 28–31 Geschäftsregeln
neg = negativtests()
regeln = [
    ("1", "Double-Booking", "Ein Exemplar kann nicht doppelt in einem überlappenden Zeitraum verliehen werden.",
     [z for z in ddl("ausleihe") if "excl_ausleihe" in z], "EXCLUDE-Constraint in 02_schema.sql (deklarativ, btree_gist)",
     "Negativ: überlappende Bestätigung 05.–12.12."),
    ("2", "Bewertung nur nach abgeschlossener eigener Ausleihe", "Eine Bewertung ist nur zu einer eigenen und abgeschlossenen Ausleihe möglich.",
     trigger_funktion("pruefe_bewertung_werk"), "Trigger BEFORE INSERT/UPDATE auf bewertung_werk (analog bewertung_person)",
     "Negativ: besitzende Person bewertet das Werk"),
    ("3", "Maximale Leihdauer", "Die maximal gesetzte Leihdauer des Exemplars kann nicht überschritten werden.",
     trigger_funktion("pruefe_leihdauer"), "Trigger BEFORE INSERT/UPDATE auf ausleihe",
     "Negativ: 15 Tage (01.–15.02.)"),
    ("4", "Zulässige Statusübergänge", "Eine Ausleihe kann nur die definierten Statusübergänge durchlaufen.",
     trigger_funktion("pruefe_statusuebergang"), "Trigger BEFORE UPDATE OF status auf ausleihe",
     "Negativ: laufend → angefragt"),
]
tests_tab = beleg("tests.txt")
for num, kurz, a6, code, wie, label in regeln:
    s = neue_folie()
    titel(s, f"Geschäftsregel {num}: {kurz}", None)
    box(s, 0.5, 0.95, 12.3, 0.5, f"A6, Regel {num}: „{a6}“", size=12, italic=True, color=GRAU)
    rtext = T.get("regeln", {}).get(num, [])
    absaetze(s, 0.5, 1.5, 6.0, 1.55, rtext, "2–3 Sätze: warum diese Umsetzung (Constraint vs. Trigger) und was der Negativtest beweist", size=11)  # 55–90 Wörter: Box 1,55 in; 10,5 pt endete bei 2,78 in (Regel 1), 11 pt geschätzt ≈ 2,9 in, Label bei 3,07
    code_box(s, 0.5, 3.35, 6.0, 3.55, code, size=7.5, label=f"Umsetzung: {wie}")
    regel_nr, stmt = neg[label]
    _, h2, _ = code_box(s, 6.8, 1.85, 6.0, 1.6, stmt + ["-- erwartet: Fehler, Zeile wird abgelehnt"], size=8.5, label=f"Negativtest ({label}, sql/init/06_tests.sql)")
    y_bild = 1.85 + h2 + 0.4
    box(s, 6.8, y_bild - 0.28, 6, 0.28, "Ergebnis (Auszug aus der Testtabelle, 06_tests.sql)", size=10, bold=True, color=GRAU)
    pfad, pw, ph = terminal_png(testtabelle_zweizeilig(tests_tab, num), f"regel_{num}", kopf=f"06_tests.sql — Fälle zu Regel {num}")
    bild_einpassen(s, pfad, pw, ph, 6.8, y_bild, 6.0, 7.0 - y_bild)
    fusszeile(s, nr, "Negativtests laufen bei jedem Containerstart; ein nicht bestandener Fall verhindert den Start (06_tests.sql)")

# 32 Testübersicht
s = neue_folie()
titel(s, "Alle 16 Regeltests auf einen Blick", T.get("tests_uebersicht"))
pfad, pw, ph = terminal_png(kompakte_testtabelle(tests_tab, breite_fall=58, breite_erg=70), "tests_uebersicht")
bild_einpassen(s, pfad, pw, ph, 0.5, 1.6, 12.3, 5.3)
fusszeile(s, nr, "Positivtests belegen, dass die Regel erlaubte Vorgänge nicht blockiert; Negativtests, dass sie verbotene abweist")

# 33–34 Performance
perf = [
    ("Umkreissuche: GiST-Index auf adresse.position", "umkreis",
     ["SELECT adresse_id, strasse, hausnummer, plz, ort",
      "  FROM adresse_perf            -- 200.000 Punkte in der Bounding-Box Frankfurt",
      " WHERE ST_DWithin(position,",
      "         ST_SetSRID(ST_MakePoint(8.6850, 50.1250), 4326)::geography,",
      "         1500);                 -- 1,5 km um den Nordend-Punkt, 14.701 Treffer"],
     "idx_adresse_position", "01a_umkreissuche_ohne_index.txt", "01b_umkreissuche_mit_index.txt"),
    ("Titelsuche: GIN-Index mit pg_trgm auf werk.titel", "trigram",
     ["SELECT werk_id, titel",
      "  FROM werk_perf               -- 200.000 generierte Titel",
      " WHERE titel ILIKE '%Turm und Spiegel%';   -- Treffer mitten im Titel, 522 Zeilen"],
     "idx_werk_titel_trgm", "02a_trigramsuche_ohne_index.txt", "02b_trigramsuche_mit_index.txt"),
]
for kopf, key, query, idxname, ohne, mit in perf:
    s = neue_folie()
    titel(s, kopf, None)
    absaetze(s, 0.5, 0.95, 9.4, 1.0, T.get("performance", {}).get(key, []), "2–3 Sätze: was gemessen wurde und was der Plan zeigt", size=12)
    t_ohne, t_mit = ausfuehrungszeit(ohne), ausfuehrungszeit(mit)
    de = lambda v: f"{v:.1f}".replace(".", ",")   # deutsches Dezimalkomma wie in den Folientexten (3,9×); Terminalbilder bleiben PostgreSQL-Format
    box(s, 10.2, 0.8, 2.7, 1.0, f"{de(t_ohne / t_mit)}×", size=40, bold=True, color=DUNKEL, align=PP_ALIGN.RIGHT)  # 40 pt: Text 0,85–1,52 in, Box 0,8–1,8 (Messung Phase 2)
    box(s, 10.2, 1.70, 2.7, 0.4, f"{de(t_ohne)} ms → {de(t_mit)} ms", size=11, color=GRAU, align=PP_ALIGN.RIGHT)  # Text ab 1,75 (Rand 0,05), endet vor 1,95
    pfad, pw, ph = terminal_png(plan_zeilen(ohne), f"perf_{key}_ohne", kopf=f"EXPLAIN (ANALYZE, BUFFERS) — ohne Index — Execution Time {t_ohne:.1f} ms")
    bild_einpassen(s, pfad, pw, ph, 0.5, 1.95, 12.3, 2.05)
    pfad, pw, ph = terminal_png(plan_zeilen(mit), f"perf_{key}_mit", kopf=f"EXPLAIN (ANALYZE, BUFFERS) — mit Index — Execution Time {t_mit:.1f} ms")
    bild_einpassen(s, pfad, pw, ph, 0.5, 4.1, 12.3, 2.02)   # endet 6,12 — Label-Text beginnt erst bei 6,15 (Box 6,10 + Rand 0,05)
    code_box(s, 0.5, 6.38, 6.0, 0.65, query, size=7.5, label="Abfrage (sql/perf/)")  # Label ab 6,10 in, frei vom Planbild (vorher 0,03 in Überlappung)
    idx = next(i for i in INDEXES if i[0] == idxname)
    code_box(s, 6.8, 6.35, 6.0, 0.3, [f"CREATE INDEX {idx[0]} ON {idx[1]} {idx[2]};"], size=7.5, label="Index (sql/init/03_indexes.sql) und Begründung aus model.py")
    box(s, 6.8, 6.62, 6.0, 0.45, idx[3], size=8, color=GRAU, italic=True)
    fusszeile(s, nr, "200.000 Demo-Zeilen (setseed 0.42), 2. Lauf, warmer Cache, jit = off · Pläne gekürzt (cost=…, Geometrie-Literal) — vollständig in sql/perf/ergebnis/")

# 35–36 Zusammenfassung der Implementierung
kz = T.get("zusammenfassung", [])
haelfte = (len(kz) + 1) // 2 if kz else 0
for i, teil in enumerate((kz[:haelfte], kz[haelfte:])):
    s = neue_folie()
    titel(s, "Zusammenfassung der Implementierung" + (" (1/2)" if i == 0 else " (2/2)"), None)
    absaetze(s, 0.6, 1.2, 12.1, 5.7, teil, "Kurzzusammenfassung, mindestens eine halbe Seite (ca. 350–450 Wörter), auf zwei Folien verteilt", size=18, abstand=14)  # 14 pt endete bei 3,9 in — halbe Folie leer
    fusszeile(s, nr, "Dieser Text wird zusätzlich als Kurzzusammenfassung (_KZ) eingereicht")

# 37 Reflexion
s = neue_folie()
titel(s, "Reflexion und Ausblick auf Phase 3", None)
aufzaehlung(s, 0.6, 1.2, 12.1, 5.6, T.get("reflexion", []), "4–5 Punkte: was schwierig war, was du anders machen würdest, was in Phase 3 folgt", size=18)
fusszeile(s, nr)

OUT_DIR.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
offen = sum(1 for b in bloecke if platzhalter(b.get("erklaerung")))
print(f"{OUT.relative_to(REPO)}: {nr} Folien · Screenshots in {SHOTS.relative_to(REPO)} · "
      f"offene @erklaerung: {offen}/21")
