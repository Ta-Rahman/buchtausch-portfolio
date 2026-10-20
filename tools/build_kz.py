#!/usr/bin/env python3
"""Kurzzusammenfassung (_KZ) der Erarbeitungsphase als Word-Datei.

Liest denselben Text wie die Folien 35/36 (tools/texte_phase2.json, Schlüssel
"zusammenfassung" und "titel") und schreibt abgabe/02_Erarbeitungsphase/
Erarbeitungsphase_KZ.docx im Layout der Phase-1-Abgabe (Arial, 2 cm Rand,
Kopfzeile Name · Matrikelnummer · Modul · Datum). Der PDF-Export erfolgt wie in
Phase 1 aus Word; zur Längenprüfung reicht `soffice --convert-to pdf`.

Aufruf: python3 tools/build_kz.py
"""
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
T = json.loads((ROOT / "tools" / "texte_phase2.json").read_text(encoding="utf-8"))
OUT = ROOT / "abgabe" / "02_Erarbeitungsphase" / "Erarbeitungsphase_KZ.docx"

SCHRIFT = "Arial"
BLAU = RGBColor(0x1F, 0x38, 0x64)
GRAU = RGBColor(0x59, 0x59, 0x59)


def absatz(doc, text, size=11, bold=False, color=None, after=6, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(0)
    if align:
        p.alignment = align
    r = p.add_run(text)
    r.font.name = SCHRIFT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), SCHRIFT)
    r.font.size = Pt(size)
    r.font.bold = bold
    if color is not None:
        r.font.color.rgb = color
    return p


def linie(p):
    """Dünne blaue Linie unter dem Absatz (wie in der Phase-1-Vorlage)."""
    pPr = p._element.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "8"), ("space", "4"), ("color", "1F3864")):
        bottom.set(qn(f"w:{k}"), v)
    bdr.append(bottom)
    pPr.append(bdr)


def main():
    kz = [a for a in T.get("zusammenfassung", []) if a.strip()]
    ti = dict(T.get("titel", {}))
    # Matrikelnummer steht nicht im Repo, sondern in der gitignorierten tools/titel_lokal.json
    lokal = ROOT / "tools" / "titel_lokal.json"
    if lokal.exists():
        ti.update(json.loads(lokal.read_text(encoding="utf-8")))
    if not kz or not all(ti.get(k) for k in ("name", "matrikel", "datum")):
        raise SystemExit("texte_phase2.json/titel_lokal.json: zusammenfassung oder titel unvollständig")

    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for seite in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, seite, Cm(2.0))
    doc.styles["Normal"].font.name = SCHRIFT
    doc.styles["Normal"].font.size = Pt(11)
    # Automatische Silbentrennung (wie für das Abstract in Phase 3 gefordert),
    # sonst reißt der Blocksatz breite Lücken.
    hy = OxmlElement("w:autoHyphenation"); hy.set(qn("w:val"), "1")
    tab = doc.settings.element.find(qn("w:defaultTabStop"))  # Schema-Reihenfolge: direkt nach defaultTabStop
    if tab is not None:
        tab.addnext(hy)
    else:
        doc.settings.element.append(hy)

    # Kopfzeile wie Phase 1: Nachname, Vorname(n) · Matrikelnummer · Modul · Datum
    teile = ti["name"].split()
    kopfname = f"{teile[-1]}, {' '.join(teile[:-1])}" if len(teile) > 1 else ti["name"]
    kopf = absatz(doc, f"{kopfname}  ·  Matrikelnummer {ti['matrikel']}  ·  DLBDSPBDM01_D  ·  {ti['datum']}",
                  size=8.5, color=GRAU, after=2)
    linie(kopf)
    absatz(doc, "Portfolio – Erarbeitungs-/Reflexionsphase", size=14, bold=True, color=BLAU, after=3)
    absatz(doc, "Data-Mart-Erstellung in SQL: Datenbank für eine Buchtausch-App", size=12, color=GRAU, after=9)
    absatz(doc, "Kurzzusammenfassung der Implementierung", size=12, bold=True, after=6)
    for a in kz:
        absatz(doc, a, size=11, after=8, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    woerter = sum(len(a.split()) for a in kz)
    print(f"geschrieben: {OUT.relative_to(ROOT)}  ({len(kz)} Absätze, {woerter} Wörter)")


if __name__ == "__main__":
    main()
